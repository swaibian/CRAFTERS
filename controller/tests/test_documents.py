"""SWE4 tests: document adapters on fixtures created in tmp_path.

Positive: xlsx round-trip (inspect/set_cell/add_sheet + verify), png
transform + rehash, pdf inspect + thumbnail + rotate.
Negative: wrong suffix, bad cell, unknown doc, bad page, delete-only-page;
originals never mutated (work on copies).
"""

from __future__ import annotations

import pytest

from controller.documents import (
    DocumentError,
    DocumentRegistry,
    PdfAdapter,
    PhotoAdapter,
    SheetAdapter,
    content_hash,
)


@pytest.fixture()
def registry(tmp_path):
    return DocumentRegistry(tmp_path / "work")


def _make_xlsx(path):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "hello"
    ws["B2"] = 42
    wb.save(path)
    wb.close()


def _make_png(path, size=(64, 48)):
    from PIL import Image

    img = Image.new("RGB", size, color=(200, 30, 30))
    img.save(path, "PNG")


def _make_pdf(path, pages=2):
    import fitz

    doc = fitz.open()
    for i in range(pages):
        page = doc.new_page(width=200, height=300)
        page.insert_text((20, 50), f"page {i + 1} text")
    doc.save(path)
    doc.close()


# --- sheets ---------------------------------------------------------------


def test_sheet_round_trip(registry, tmp_path):
    src = tmp_path / "book.xlsx"
    _make_xlsx(src)
    before_hash = content_hash(src)
    adapter = SheetAdapter(registry)
    doc = adapter.open(src, "book.xlsx")
    assert doc.work_path != src
    assert content_hash(src) == before_hash  # original untouched

    info = adapter.inspect(doc.doc_id)
    assert info["sheets"][0]["title"] == "Data"
    assert info["sheets"][0]["max_row"] >= 1

    evidence = adapter.set_cell(doc.doc_id, "A1", "world")
    assert evidence["before"] == "hello"
    assert evidence["after"] == "world"
    assert evidence["content_hash"] != info["content_hash"]

    added = adapter.add_sheet(doc.doc_id, "Second")
    assert "Second" in added["after_sheets"]
    assert len(added["after_sheets"]) == len(added["before_sheets"]) + 1


def test_sheet_negative(registry, tmp_path):
    adapter = SheetAdapter(registry)
    bad = tmp_path / "note.txt"
    bad.write_text("nope")
    with pytest.raises(DocumentError):
        adapter.open(bad, "note.txt")
    src = tmp_path / "book.xlsx"
    _make_xlsx(src)
    doc = adapter.open(src, "book.xlsx")
    with pytest.raises(DocumentError):
        adapter.set_cell(doc.doc_id, "not-a-cell", "x")
    with pytest.raises(DocumentError):
        adapter.add_sheet(doc.doc_id, "")
    with pytest.raises(DocumentError):
        adapter.inspect("doc_nope")


# --- photos ---------------------------------------------------------------


def test_photo_transform_and_rehash(registry, tmp_path):
    src = tmp_path / "pic.png"
    _make_png(src)
    before_hash = content_hash(src)
    adapter = PhotoAdapter(registry)
    doc = adapter.open(src, "pic.png")

    info = adapter.inspect(doc.doc_id)
    assert (info["width"], info["height"]) == (64, 48)
    assert info["mode"] == "RGB"

    resized = adapter.transform(doc.doc_id, "resize", width=800)
    assert resized["after"]["width"] == 800
    assert resized["after"]["content_hash"] != info["content_hash"]

    gray = adapter.transform(doc.doc_id, "grayscale")
    assert gray["after"]["mode"] == "L"

    rotated = adapter.transform(doc.doc_id, "rotate", degrees=90)
    assert (rotated["after"]["width"], rotated["after"]["height"]) == (
        rotated["before"]["height"],
        rotated["before"]["width"],
    )
    assert content_hash(src) == before_hash  # original untouched


def test_photo_negative(registry, tmp_path):
    adapter = PhotoAdapter(registry)
    bad = tmp_path / "doc.pdf"
    bad.write_bytes(b"%PDF-fake")
    with pytest.raises(DocumentError):
        adapter.open(bad, "doc.pdf")
    src = tmp_path / "pic.png"
    _make_png(src)
    doc = adapter.open(src, "pic.png")
    with pytest.raises(DocumentError):
        adapter.transform(doc.doc_id, "resize")
    with pytest.raises(DocumentError):
        adapter.transform(doc.doc_id, "frobnicator")


# --- pdfs -----------------------------------------------------------------


def test_pdf_inspect_thumbnail_rotate(registry, tmp_path):
    src = tmp_path / "doc.pdf"
    _make_pdf(src)
    before_hash = content_hash(src)
    adapter = PdfAdapter(registry)
    doc = adapter.open(src, "doc.pdf")

    info = adapter.inspect(doc.doc_id)
    assert len(info["pages"]) == 2
    assert info["pages"][0]["text_chars"] > 0

    thumb = adapter.render_thumbnail(doc.doc_id, page=0)
    assert thumb.exists() and thumb.stat().st_size > 0

    rotated = adapter.transform(doc.doc_id, "rotate", page=0, degrees=90)
    assert rotated["after"]["pages"][0]["rotation"] == 90
    assert rotated["after"]["content_hash"] != info["content_hash"]
    assert content_hash(src) == before_hash  # original untouched


def test_pdf_delete_and_negative(registry, tmp_path):
    adapter = PdfAdapter(registry)
    src = tmp_path / "doc.pdf"
    _make_pdf(src, pages=2)
    doc = adapter.open(src, "doc.pdf")
    deleted = adapter.transform(doc.doc_id, "delete", page=0)
    assert len(deleted["after"]["pages"]) == 1

    single = tmp_path / "one.pdf"
    _make_pdf(single, pages=1)
    doc1 = adapter.open(single, "one.pdf")
    with pytest.raises(DocumentError):
        adapter.transform(doc1.doc_id, "delete", page=0)
    with pytest.raises(DocumentError):
        adapter.transform(doc.doc_id, "rotate", page=99)
    with pytest.raises(DocumentError):
        adapter.render_thumbnail(doc.doc_id, page=99)
